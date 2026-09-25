#!/usr/bin/env python3
"""Canonical resume renderer: modern-sidebar-v1.

Input: JSON file containing already-approved resume content.
Output: DOCX. PDF conversion is a separate deterministic step.

This renderer owns visual presentation only. It must not invent or rewrite
candidate facts. Content selection happens upstream from Candidate Knowledge.
"""
import argparse, json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

NAVY="102A43"; BLUE="345B7E"; SIDEBAR="EAF2F8"; TEXT="273444"
RENDERER_KEY="modern-sidebar-v1"

def fmt(run,size,bold=False,color=TEXT):
    run.font.name="Arial"; run.font.size=Pt(size); run.font.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)

def cell_margins(cell, top=80,start=100,bottom=70,end=100):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar=OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for k,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        el=tcMar.find(qn("w:"+k))
        if el is None: el=OxmlElement("w:"+k); tcMar.append(el)
        el.set(qn("w:w"),str(v)); el.set(qn("w:type"),"dxa")

def no_borders(table):
    b=OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement("w:"+edge); e.set(qn("w:val"),"nil"); b.append(e)
    table._tbl.tblPr.append(b)

def rule(p):
    pb=OxmlElement("w:pBdr"); bt=OxmlElement("w:bottom")
    for k,v in (("val","single"),("sz","5"),("space","2"),("color",BLUE)):
        bt.set(qn("w:"+k),v)
    pb.append(bt); p._p.get_or_add_pPr().append(pb)

def heading(cell,text,sidebar=False):
    p=cell.add_paragraph(); p.paragraph_format.space_before=Pt(4 if sidebar else 3); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(text.upper()); fmt(r,9.1 if sidebar else 10,True,NAVY); rule(p)

def bullet(cell,text,sidebar=False):
    p=cell.add_paragraph(); p.paragraph_format.left_indent=Inches(.12 if sidebar else .16)
    p.paragraph_format.first_line_indent=Inches(-.09); p.paragraph_format.space_after=Pt(.45); p.paragraph_format.line_spacing=1
    r=p.add_run("• "); fmt(r,7.55 if sidebar else 7.85,True,NAVY)
    r=p.add_run(text); fmt(r,7.55 if sidebar else 7.85,False,TEXT)

def add_fixed_sidebar(section):
    # Page-level VML shape in the header, behind document text.
    header=section.header
    p=header.paragraphs[0]
    pict=OxmlElement("w:pict")
    shape=parse_xml(
        '<v:rect %s style="position:absolute;left:0;top:0;width:2.42in;height:11in;'
        'z-index:-251654144;mso-position-horizontal-relative:page;'
        'mso-position-vertical-relative:page" fillcolor="#%s" stroked="f"/>'
        % (nsdecls("v"), SIDEBAR)
    )
    pict.append(shape); p._p.append(pict)
    header.is_linked_to_previous=False
    section.header_distance=Inches(0)

def role(cell,item):
    p=cell.add_paragraph(); p.paragraph_format.space_before=Pt(2.5); p.paragraph_format.space_after=Pt(1)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(5.05), WD_TAB_ALIGNMENT.RIGHT)
    r=p.add_run(item["title"]); fmt(r,8.75,True,NAVY)
    r=p.add_run("  |  "+item["company"]); fmt(r,8.15,False,TEXT)
    r=p.add_run("\t"+item["dates"]); fmt(r,7.65,False,NAVY)
    for x in item.get("bullets",[]): bullet(cell,x,False)

def render(data,out_path):
    d=Document(); s=d.sections[0]
    s.top_margin=Inches(.26); s.bottom_margin=Inches(.26); s.left_margin=Inches(.28); s.right_margin=Inches(.28)
    add_fixed_sidebar(s)
    d.styles["Normal"].font.name="Arial"; d.styles["Normal"].font.size=Pt(8.3); d.styles["Normal"].font.color.rgb=RGBColor.from_string(TEXT)

    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(data["name"]); fmt(r,27,True,"0B2239")
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
    r=p.add_run(data["headline"]); fmt(r,9,False,NAVY)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(3)
    r=p.add_run(data["contact_line"]); fmt(r,7.35,False,TEXT)

    t=d.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    t.columns[0].width=Inches(2.12); t.columns[1].width=Inches(5.72)
    left,right=t.rows[0].cells; no_borders(t)
    for c in (left,right): c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; c.paragraphs[0].paragraph_format.space_after=Pt(0)
    # Transparent sidebar cell. The page-level background provides the blue.
    cell_margins(left,105,120,80,120); cell_margins(right,105,150,80,85)

    for sec in data.get("sidebar",[]):
        heading(left,sec["title"],True)
        for x in sec.get("items",[]): bullet(left,x,True)

    heading(right,"Professional Summary")
    p=right.add_paragraph(); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1
    r=p.add_run(data["summary"]); fmt(r,8.2,False,TEXT)

    heading(right,"Professional Experience")
    for item in data.get("experience",[]): role(right,item)

    heading(right,"Education")
    for edu in data.get("education",[]):
        p=right.add_paragraph(); p.paragraph_format.space_after=Pt(0)
        r=p.add_run(edu["degree"]); fmt(r,8.25,True,NAVY)
        r=p.add_run("  |  "+edu["school"]); fmt(r,8.1,False,TEXT)

    d.core_properties.title=f'{data["name"]} Resume'
    d.core_properties.subject=f'Renderer: {RENDERER_KEY}'
    d.save(out_path)

if __name__=="__main__":
    ap=argparse.ArgumentParser()
    ap.add_argument("input_json"); ap.add_argument("output_docx")
    args=ap.parse_args()
    render(json.loads(Path(args.input_json).read_text()), args.output_docx)
