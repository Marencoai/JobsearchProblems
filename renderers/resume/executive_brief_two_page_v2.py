from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import argparse, json

NAVY='15324B'; BLUE='4A7396'; PALE='EEF4F8'; PALE2='F7FAFC'
TEXT='263746'; MUTED='66727D'; RULE='8CA8BC'; FONT='Arial'
RENDERER_KEY='executive-brief-two-page-v2'

def set_run(run,size,bold=False,color=TEXT,italic=False):
    run.font.name=FONT
    run.font.size=Pt(size)
    run.font.bold=bold
    run.font.italic=italic
    run.font.color.rgb=RGBColor.from_string(color)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=tcPr.find(qn('w:shd'))
    if shd is None:
        shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=70,start=85,bottom=70,end=85):
    tcPr=cell._tc.get_or_add_tcPr()
    tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None:
            node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def no_borders(table):
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge)
        el.set(qn('w:val'),'nil')
        borders.append(el)
    table._tbl.tblPr.append(borders)

def border_cell(cell, color=RULE, sz='5'):
    tcPr=cell._tc.get_or_add_tcPr()
    borders=tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders=OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right'):
        el=OxmlElement('w:'+edge)
        el.set(qn('w:val'),'single')
        el.set(qn('w:sz'),sz)
        el.set(qn('w:color'),color)
        borders.append(el)

def add_rule(p, color=RULE, sz='6'):
    pBdr=OxmlElement('w:pBdr')
    bottom=OxmlElement('w:bottom')
    bottom.set(qn('w:val'),'single')
    bottom.set(qn('w:sz'),sz)
    bottom.set(qn('w:space'),'2')
    bottom.set(qn('w:color'),color)
    pBdr.append(bottom)
    p._p.get_or_add_pPr().append(pBdr)

def section_heading(doc,text, before=7, after=4):
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(before)
    p.paragraph_format.space_after=Pt(after)
    set_run(p.add_run(text.upper()),10.3,True,NAVY)
    add_rule(p)

def add_header(doc,data,compact=False):
    if not compact:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(0)
        set_run(p.add_run(data['name']),26,True,NAVY)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(1)
        set_run(p.add_run(data['headline']),9.8,True,BLUE)
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(5)
        set_run(p.add_run(data['contact']),7.9,False,MUTED)
        add_rule(p,BLUE,'7')
    else:
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1)
        set_run(p.add_run(data['name']),14.2,True,NAVY)
        set_run(p.add_run('  |  '+data['headline']),8.3,False,MUTED)
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
        set_run(p.add_run(data['contact']),7.5,False,MUTED)
        add_rule(p,BLUE,'5')

def add_summary(doc,text):
    section_heading(doc,'Professional Summary',7,4)
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.line_spacing=1.16
    set_run(p.add_run(text),10.0,False,TEXT)

def add_capabilities(doc,items):
    section_heading(doc,'Core Capabilities',7,4)
    table=doc.add_table(rows=0, cols=3)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    no_borders(table)
    widths=[2.42,2.42,2.42]
    for start in range(0,len(items),3):
        cells=table.add_row().cells
        for j,c in enumerate(cells):
            c.width=Inches(widths[j])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c,95,75,95,75)
            if start+j < len(items):
                shade(c,PALE2)
                p=c.paragraphs[0]
                p.paragraph_format.space_after=Pt(0)
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                set_run(p.add_run(items[start+j]),8.8,True,NAVY)

def add_impact(doc,items):
    section_heading(doc,'Selected Impact',8,4)
    table=doc.add_table(rows=2, cols=2)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    no_borders(table)
    idx=0
    for r in range(2):
        for c in range(2):
            cell=table.cell(r,c)
            cell.width=Inches(3.62)
            set_cell_margins(cell,125,100,125,100)
            shade(cell,PALE)
            if idx < len(items):
                metric,label=items[idx]
                p=cell.paragraphs[0]
                p.paragraph_format.space_after=Pt(1)
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                set_run(p.add_run(metric),14.0,True,BLUE)
                p=cell.add_paragraph()
                p.paragraph_format.space_after=Pt(0)
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                set_run(p.add_run(label),8.75,False,TEXT)
            idx+=1

def add_projects(doc,projects):
    section_heading(doc,'Selected Projects & Systems',8,4)
    table=doc.add_table(rows=0, cols=1)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    no_borders(table)
    for pr in projects:
        cell=table.add_row().cells[0]
        set_cell_margins(cell,125,110,125,110)
        shade(cell,PALE2)
        border_cell(cell,'D9E5ED','3')
        p=cell.paragraphs[0]
        p.paragraph_format.space_after=Pt(1)
        set_run(p.add_run(pr['name']),9.2,True,NAVY)
        if pr.get('role'):
            set_run(p.add_run('  |  '+pr['role']),8.3,False,MUTED)
        p=cell.add_paragraph()
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1.04
        set_run(p.add_run(pr['text']),8.9,False,TEXT)

def add_supporting(doc,tools,certs,education):
    section_heading(doc,'Supporting Information',8,4)
    table=doc.add_table(rows=1,cols=2)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    no_borders(table)
    table.columns[0].width=Inches(4.65)
    table.columns[1].width=Inches(2.6)
    left,right=table.rows[0].cells
    for cell in (left,right):
        set_cell_margins(cell,125,100,125,100)
        shade(cell,PALE)
    p=left.paragraphs[0]
    p.paragraph_format.space_after=Pt(2)
    set_run(p.add_run('TOOLS & PLATFORMS'),8.7,True,NAVY)
    p=left.add_paragraph()
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=1.03
    set_run(p.add_run('  |  '.join(tools)),8.55,False,TEXT)
    p=right.paragraphs[0]
    p.paragraph_format.space_after=Pt(2)
    set_run(p.add_run('CREDENTIALS'),8.7,True,NAVY)
    for x in certs+education:
        p=right.add_paragraph()
        p.paragraph_format.space_after=Pt(1)
        set_run(p.add_run(x),8.45,False,TEXT)

def add_role(doc,role):
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(7)
    p.paragraph_format.space_after=Pt(1.5)
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.85),WD_TAB_ALIGNMENT.RIGHT)
    set_run(p.add_run(role['title']),9.85,True,NAVY)
    set_run(p.add_run('  |  '+role['company']),9.25,False,TEXT)
    set_run(p.add_run('\t'+role['dates']),8.6,False,MUTED)
    if role.get('scope'):
        sp=doc.add_paragraph()
        sp.paragraph_format.space_after=Pt(3.1)
        sp.paragraph_format.left_indent=Inches(.02)
        sp.paragraph_format.line_spacing=1.04
        set_run(sp.add_run(role['scope']),8.85,False,MUTED,True)
    for b in role['bullets']:
        bp=doc.add_paragraph()
        bp.paragraph_format.left_indent=Inches(.18)
        bp.paragraph_format.first_line_indent=Inches(-.12)
        bp.paragraph_format.space_after=Pt(2.9)
        bp.paragraph_format.line_spacing=1.08
        bp.paragraph_format.keep_together=True
        set_run(bp.add_run('• '),8.2,True,BLUE)
        set_run(bp.add_run(b),9.05,False,TEXT)

def render(data,out):
    d=Document()
    sec=d.sections[0]
    sec.page_width=Inches(8.5)
    sec.page_height=Inches(11)
    sec.top_margin=Inches(.42)
    sec.bottom_margin=Inches(.38)
    sec.left_margin=Inches(.58)
    sec.right_margin=Inches(.58)
    d.styles['Normal'].font.name=FONT
    d.styles['Normal'].font.size=Pt(8.5)

    add_header(d,data,False)
    add_summary(d,data['summary'])
    add_capabilities(d,data['capabilities'])
    add_impact(d,data['impact'])
    add_projects(d,data['projects'])
    add_supporting(d,data['tools'],data['certifications'],data['education'])

    d.add_page_break()
    add_header(d,data,True)
    section_heading(d,'Professional Experience',4,3)
    for role in data['experience']:
        add_role(d,role)

    d.core_properties.title=f"{data['name']} - {data['headline']}"
    d.core_properties.subject=f'Renderer: {RENDERER_KEY}'
    d.save(out)

if __name__=="__main__":
    ap=argparse.ArgumentParser()
    ap.add_argument("input_json")
    ap.add_argument("output_docx")
    args=ap.parse_args()
    render(json.loads(Path(args.input_json).read_text()), args.output_docx)
